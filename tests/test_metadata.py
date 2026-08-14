import pytest
import openpyxl

from src.chunking.baseline import baseline_chunk_articles
from src.chunking.structure_aware import structure_aware_chunk_articles
from src.config import MARKDOWN_ARTICLES_DIR
from src.ingestion.loader import load_article, load_articles
from src.ingestion.metadata import validate_metadata

REQUIRED_CHUNK_FIELDS = {
    "source_file",
    "article_id",
    "product_area",
    "last_updated",
    "chunk_id",
}


def test_all_six_articles_load():
    articles = load_articles()
    assert len(articles) == 6


def test_missing_metadata_raises():
    with pytest.raises(ValueError):
        validate_metadata({"source_file": "x.md", "article_id": "HC-1"})


def test_baseline_chunks_have_required_metadata():
    articles = load_articles()
    chunks = baseline_chunk_articles(articles)

    assert chunks
    for chunk in chunks:
        assert REQUIRED_CHUNK_FIELDS.issubset(chunk.metadata.keys())


def test_structure_aware_chunks_have_required_metadata():
    # Structure-aware chunker requires ## section headings — use the markdown articles.
    articles = load_articles(MARKDOWN_ARTICLES_DIR)
    chunks = structure_aware_chunk_articles(articles)

    assert chunks
    for chunk in chunks:
        assert REQUIRED_CHUNK_FIELDS.issubset(chunk.metadata.keys())


def test_load_article_supports_plain_text_without_front_matter(tmp_path):
    file_path = tmp_path / "notes.txt"
    file_path.write_text("Plain text content", encoding="utf-8")

    article = load_article(file_path)

    assert article.content == "Plain text content"
    assert article.metadata["source_file"] == file_path.name
    assert article.metadata["article_id"] == "notes"
    assert article.metadata["product_area"] == "unknown"


def test_load_article_supports_xlsx(tmp_path):
    file_path = tmp_path / "report.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fees"
    ws.append(["Plan", "Monthly Fee", "Annual Fee"])
    ws.append(["Basic", "9.99", "99.00"])
    ws.append(["Pro", "29.99", "299.00"])
    wb.save(file_path)

    article = load_article(file_path)

    assert "Fees" in article.content
    assert "Plan" in article.content
    assert "Basic" in article.content
    assert "29.99" in article.content
    assert article.metadata["article_id"] == "report"


def test_xlsx_produces_markdown_table(tmp_path):
    file_path = tmp_path / "errors.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Error Codes"
    ws.append(["Code", "Description"])
    ws.append(["ERR-001", "Invalid payment"])
    wb.save(file_path)

    article = load_article(file_path)

    assert "| Code | Description |" in article.content
    assert "| ERR-001 | Invalid payment |" in article.content


def test_xlsx_skips_empty_sheets(tmp_path):
    file_path = tmp_path / "mixed.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Data"
    ws1.append(["Name", "Value"])
    ws1.append(["Alpha", "1"])
    wb.create_sheet("Empty")
    wb.save(file_path)

    article = load_article(file_path)

    assert "Data" in article.content
    assert "Empty" not in article.content
