from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml

try:
    import docx  # type: ignore
except Exception:
    docx = None

try:
    from pypdf import PdfReader as _PdfReader  # type: ignore
except Exception:
    _PdfReader = None

try:
    import openpyxl  # type: ignore
except Exception:
    openpyxl = None

try:
    import xlrd  # type: ignore
except Exception:
    xlrd = None

from src.config import ARTICLES_DIR
from src.ingestion.metadata import validate_metadata


@dataclass
class Article:
    content: str
    metadata: dict[str, Any]


def _split_front_matter(raw_text: str) -> tuple[dict[str, Any], str]:
    if not raw_text.startswith("---"):
        return {}, raw_text

    _, front_matter_block, body = raw_text.split("---", 2)
    front_matter = yaml.safe_load(front_matter_block) or {}
    return front_matter, body.strip()


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").strip()


def _extract_xlsx(file_path: Path) -> str:
    if openpyxl is None:
        raise ValueError("openpyxl is required for .xlsx files: pip install openpyxl")

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sections: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [
            [_cell_str(cell) for cell in row]
            for row in ws.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue

        headers = rows[0]
        data_rows = rows[1:]

        lines = [f"## {sheet_name}", ""]
        if headers:
            lines.append("| " + " | ".join(headers) + " |")
            lines.append("| " + " | ".join("---" for _ in headers) + " |")
            for row in data_rows:
                padded = row + [""] * (len(headers) - len(row))
                lines.append("| " + " | ".join(padded[: len(headers)]) + " |")

        sections.append("\n".join(lines))

    wb.close()
    return "\n\n".join(sections)


def _extract_xls(file_path: Path) -> str:
    if xlrd is None:
        raise ValueError("xlrd is required for .xls files: pip install xlrd")

    wb = xlrd.open_workbook(str(file_path))
    sections: list[str] = []

    for sheet in wb.sheets():
        if sheet.nrows == 0:
            continue

        headers = [_cell_str(sheet.cell_value(0, col)) for col in range(sheet.ncols)]
        lines = [f"## {sheet.name}", ""]
        if headers:
            lines.append("| " + " | ".join(headers) + " |")
            lines.append("| " + " | ".join("---" for _ in headers) + " |")
            for row_idx in range(1, sheet.nrows):
                row = [_cell_str(sheet.cell_value(row_idx, col)) for col in range(sheet.ncols)]
                if any(row):
                    lines.append("| " + " | ".join(row) + " |")

        sections.append("\n".join(lines))

    return "\n\n".join(sections)


def _extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix in (".md", ".txt"):
        return file_path.read_text(encoding="utf-8")

    if suffix == ".docx":
        if docx is None:
            raise ValueError("python-docx is required for .docx files: pip install python-docx")
        document = docx.Document(str(file_path))
        return "\n".join(p.text for p in document.paragraphs if p.text)

    if suffix == ".pdf":
        if _PdfReader is None:
            raise ValueError("pypdf is required for .pdf files: pip install pypdf")
        try:
            reader = _PdfReader(str(file_path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            raise ValueError(
                f"Failed to extract text from PDF {file_path.name}: {e}. "
                "The PDF may be corrupted or have encoding issues."
            ) from e

    if suffix == ".xlsx":
        return _extract_xlsx(file_path)

    if suffix == ".xls":
        return _extract_xls(file_path)

    raise ValueError(f"Unsupported file type: {file_path.suffix}")


def load_article(file_path: Path) -> Article:
    raw_text = _extract_text(file_path)
    front_matter, body = _split_front_matter(raw_text)

    metadata = dict(front_matter)
    metadata.setdefault("article_id", file_path.stem)
    metadata.setdefault("product_area", "unknown")
    metadata.setdefault("last_updated", "unknown")
    metadata["source_file"] = file_path.name
    metadata["last_updated"] = str(metadata.get("last_updated", "unknown"))

    validate_metadata(metadata)

    return Article(content=body, metadata=metadata)


def load_articles(articles_dir: Path = ARTICLES_DIR) -> list[Article]:
    supported_extensions = (".md", ".txt", ".docx", ".pdf", ".xlsx", ".xls")
    file_paths = sorted(
        path
        for path in articles_dir.iterdir()
        if path.is_file() and path.suffix.lower() in supported_extensions
    )
    return [load_article(file_path) for file_path in file_paths]
